# -*- coding: utf-8 -*-
"""
Export hồ sơ BDC thành .docx hoàn chỉnh (template + user fills).

Usage: python export.py <hoso.json> [--out output.docx]

hoso.json gồm:
  { id, ten, maCongTrinhFull, heThongIds: [...],
    thongTinDuAn: {tenCongTrinh, diaDiem, chuDauTu, diaChiDaiDien, donViTuVan, canBoThamDinh},
    verdict: {"bdcId|si|subi|ii": "pass"|"kn"|"na"},
    thietKe: {"bdcId|si|subi|ii": "..."} }
"""
import json, sys, os, argparse, tempfile
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from docxcompose.composer import Composer

RED   = RGBColor(0xC0, 0x00, 0x00)
GREEN_FILL  = '92D050'
YELLOW_FILL = 'FFFF00'


def shade_cell(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, 'source')
DATA = os.path.join(ROOT, 'data')

PREAMBLE_LABELS = [
    ('tenCongTrinh',  '1. Tên công trình:'),
    ('diaDiem',       '2. Địa điểm xây dựng:'),
    ('chuDauTu',      '3. Chủ đầu tư:'),
    ('diaChiDaiDien', '4. Địa chỉ của đại diện'),
    ('donViTuVan',    '5. Đơn vị tư vấn'),
    ('canBoThamDinh', '6. Cán bộ thẩm định:'),
]

VERDICT_TEXT = {'pass': 'Đạt', 'kn': 'Không đạt', 'na': 'N/A'}
VERDICT_FILL = {'pass': GREEN_FILL, 'kn': YELLOW_FILL}


def load_parsed(bdc_id):
    path = os.path.join(DATA, f'checklist_{bdc_id}.js')
    with open(path, encoding='utf-8') as f:
        txt = f.read()
    return json.loads(txt.split('=', 1)[1].rstrip(';\n '))


def set_cell_text(cell, text, color=None):
    """Ghi text vào cell, giữ định dạng của run đầu tiên, xóa runs/paras khác."""
    if not text:
        return
    if cell.paragraphs:
        p = cell.paragraphs[0]
        if p.runs:
            for r in p.runs[1:]:
                r.text = ''
            p.runs[0].text = text
            run = p.runs[0]
        else:
            run = p.add_run(text)
        for extra_p in cell.paragraphs[1:]:
            extra_p._element.getparent().remove(extra_p._element)
    else:
        p = cell.add_paragraph(text)
        run = p.runs[0]
    if color is not None:
        run.font.color.rgb = color


def fill_preamble(doc, thong_tin):
    """Tìm các đoạn '1. Tên công trình:', '2. Địa điểm:'... và ghi giá trị user."""
    if not thong_tin:
        return
    for p in doc.paragraphs:
        txt = p.text.strip()
        for key, label in PREAMBLE_LABELS:
            if txt.startswith(label):
                val = (thong_tin.get(key) or '').strip()
                if val:
                    new = f'{label} {val}'
                    if p.runs:
                        for r in p.runs[1:]:
                            r.text = ''
                        p.runs[0].text = new
                    else:
                        p.add_run(new)
                break


def fill_main_table(doc, parsed, bdc_id, verdict, thietKe):
    """Tìm bảng BDC chính và fill cột 2 (thiết kế) + cột 5 (kết luận)."""
    # Find main table by header
    main = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip().upper() == 'TT':
            if main is None or len(t.rows) > len(main.rows):
                main = t
    if main is None:
        return 0

    n_filled = 0
    for si, sec in enumerate(parsed):
        for subi, sub in enumerate(sec['subs']):
            for ii, it in enumerate(sub['items']):
                row_idx = it.get('_row')
                if row_idx is None or row_idx >= len(main.rows):
                    continue
                key = f'{bdc_id}|{si}|{subi}|{ii}'
                v = verdict.get(key, '')
                tk = thietKe.get(key, '')
                if not v and not tk:
                    continue
                row = main.rows[row_idx]
                cells = row.cells
                if tk and len(cells) > 2:
                    set_cell_text(cells[2], tk, color=RED)
                if v and len(cells) > 5:
                    txt = VERDICT_TEXT.get(v, '')
                    if txt:
                        set_cell_text(cells[5], txt)
                    fill = VERDICT_FILL.get(v)
                    if fill:
                        shade_cell(cells[5], fill)
                n_filled += 1
    return n_filled


def fill_one(bdc_id, hoso, out_path):
    src_path = os.path.join(SRC, f'{bdc_id}.docx')
    doc = Document(src_path)
    parsed = load_parsed(bdc_id)
    if bdc_id == hoso.get('maCongTrinhFull'):
        fill_preamble(doc, hoso.get('thongTinDuAn', {}))
    n = fill_main_table(doc, parsed, bdc_id, hoso.get('verdict', {}), hoso.get('thietKe', {}))
    doc.save(out_path)
    return n


def docx_to_pdf(docx_path, pdf_path):
    """Dùng Word COM trên Windows để convert .docx → .pdf (wdFormatPDF = 17)."""
    import win32com.client
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True)
        doc.SaveAs2(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close(False)
    finally:
        word.Quit()


def export(hoso_path, out_path=None, pdf=False):
    with open(hoso_path, encoding='utf-8') as f:
        hoso = json.load(f)
    if not out_path:
        out_path = os.path.join(os.path.dirname(hoso_path) or '.', f'hoso_{hoso["id"]}.docx')

    bdc_ids = []
    if hoso.get('maCongTrinhFull'):
        bdc_ids.append(hoso['maCongTrinhFull'])
    bdc_ids += hoso.get('heThongIds', [])

    if not bdc_ids:
        raise ValueError('Hồ sơ không có BDC nào.')

    tmp_dir = tempfile.mkdtemp(prefix='bdc_export_')
    tmp_files = []
    total_filled = 0
    for bid in bdc_ids:
        tmp = os.path.join(tmp_dir, f'{bid}.docx')
        n = fill_one(bid, hoso, tmp)
        total_filled += n
        tmp_files.append(tmp)
        print(f'  filled {bid}: {n} mục')

    # Merge
    master = Document(tmp_files[0])
    if len(tmp_files) > 1:
        composer = Composer(master)
        for f in tmp_files[1:]:
            composer.append(Document(f))
        composer.save(out_path)
    else:
        master.save(out_path)

    print(f'\nOK: {out_path}')
    print(f'  Tổng số mục đã điền: {total_filled}')
    print(f'  Tổng số BDC: {len(bdc_ids)}')

    if pdf:
        pdf_path = os.path.splitext(out_path)[0] + '.pdf'
        print(f'\nĐang convert sang PDF...')
        docx_to_pdf(out_path, pdf_path)
        print(f'OK: {pdf_path}')
        return out_path, pdf_path
    return out_path


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('hoso_json')
    ap.add_argument('--out', default=None)
    ap.add_argument('--pdf', action='store_true', help='Convert thêm sang PDF (cần MS Word trên Windows)')
    args = ap.parse_args()
    export(args.hoso_json, args.out, pdf=args.pdf)
