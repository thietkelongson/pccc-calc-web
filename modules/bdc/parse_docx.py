# -*- coding: utf-8 -*-
"""
Parse 1 file .docx BDC -> JSON 3 cap + trich xuat hinh anh.
Usage: python parse_docx.py <file_id>   (vd: a4_van_phong, b1_bao_chay_thuong)
"""
import sys, os, json, re
from docx import Document
from docx.oxml.ns import qn

SRC = os.path.join(os.path.dirname(__file__), 'source')
OUT = os.path.join(os.path.dirname(__file__), 'data')

def clean(s):
    s = (s or '').replace('\r', '').strip()
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r'\n{2,}', '\n', s)
    return s

def extract_cell(cell, doc_part, file_id, counter, img_outdir):
    """Tra ve text cua cell, chen marker [[IMG:fname]] o cho co anh."""
    lines = []
    for para in cell.paragraphs:
        line_parts = []
        for run in para.runs:
            t = run.text
            if t:
                line_parts.append(t)
            # tim drawing trong run
            for drawing in run._element.findall('.//' + qn('w:drawing')):
                blip = drawing.find('.//' + qn('a:blip'))
                if blip is None:
                    continue
                rId = blip.get(qn('r:embed'))
                if not rId:
                    continue
                try:
                    image_part = doc_part.related_parts[rId]
                    ext = (image_part.partname.rsplit('.', 1)[-1] or 'png').lower()
                    if ext == 'jpeg':
                        ext = 'jpg'
                    counter[0] += 1
                    fname = f'img{counter[0]:02d}.{ext}'
                    os.makedirs(img_outdir, exist_ok=True)
                    with open(os.path.join(img_outdir, fname), 'wb') as f:
                        f.write(image_part.blob)
                    line_parts.append(f'[[IMG:{fname}]]')
                except Exception as e:
                    line_parts.append(f'[[IMG:error]]')
        lines.append(''.join(line_parts))
    return clean('\n'.join(lines))

def find_bdc_table(doc):
    """Tim bang chinh: dong dau co 'TT' va >= 5 cot. Chon bang lon nhat trong cac bang khop."""
    candidates = []
    for t in doc.tables:
        if not t.rows: continue
        first = [c.text.strip() for c in t.rows[0].cells]
        if first and first[0].upper() == 'TT' and len(first) >= 5:
            candidates.append(t)
    if candidates:
        return max(candidates, key=lambda t: len(t.rows))
    return max(doc.tables, key=lambda t: len(t.rows))

def parse(file_id):
    path = os.path.join(SRC, file_id + '.docx')
    d = Document(path)
    table = find_bdc_table(d)

    counter = [0]  # global image counter
    img_outdir = os.path.join(OUT, 'img', file_id)
    # Clear old images first
    if os.path.isdir(img_outdir):
        for f in os.listdir(img_outdir):
            try: os.remove(os.path.join(img_outdir, f))
            except: pass

    rows = []
    for ri, row in enumerate(table.rows):
        cell_texts = [extract_cell(c, d.part, file_id, counter, img_outdir) for c in row.cells]
        if cell_texts[0] == 'TT':
            continue
        tt = cell_texts[0]
        noi_dung = cell_texts[1]
        huong_dan = cell_texts[2] if len(cell_texts) > 2 else ''
        quy_dinh = cell_texts[3] if len(cell_texts) > 3 else ''
        vien_dan = cell_texts[4] if len(cell_texts) > 4 else ''
        if huong_dan == noi_dung:
            huong_dan = ''
        rows.append({'_row': ri, 'tt': tt, 'ten': noi_dung, 'huong_dan': huong_dan, 'quy_dinh': quy_dinh, 'vien_dan': vien_dan})

    sections = []
    cur_super = None  # Roman numeral I/II/III/...
    cur_sec = None
    cur_sub = None
    for r in rows:
        tt = r['tt']
        # Super-section: Roman numerals (I, II, III, IV, V, ...)
        if re.fullmatch(r'[IVX]+', tt):
            cur_super = tt
            cur_sec = None; cur_sub = None
            continue
        if re.fullmatch(r'\d+', tt):
            sec_tt = f'{cur_super}.{tt}' if cur_super else tt
            cur_sec = {
                '_row': r['_row'],
                'tt': sec_tt, 'ten': r['ten'], 'subs': [],
                'huong_dan': r['huong_dan'], 'quy_dinh': r['quy_dinh'], 'vien_dan': r['vien_dan']
            }
            sections.append(cur_sec)
            cur_sub = None
        elif re.fullmatch(r'\d+(\.\d+)+', tt):
            sub_tt = f'{cur_super}.{tt}' if cur_super else tt
            cur_sub = {
                '_row': r['_row'],
                'tt': sub_tt, 'ten': r['ten'],
                'huong_dan': r['huong_dan'],
                'quy_dinh': r['quy_dinh'], 'vien_dan': r['vien_dan'],
                'items': []
            }
            if cur_sec is None:
                cur_sec = {'tt': '?', 'ten': '(khong xac dinh)', 'subs': [], 'huong_dan':'', 'quy_dinh':'', 'vien_dan':''}
                sections.append(cur_sec)
            cur_sec['subs'].append(cur_sub)
        elif tt in ('-', '+', ''):
            item = {
                '_row': r['_row'],
                'bullet': tt or '-',
                'ten': r['ten'],
                'huong_dan': r['huong_dan'],
                'quy_dinh': r['quy_dinh'],
                'vien_dan': r['vien_dan']
            }
            if cur_sub is None:
                cur_sub = {'tt': '', 'ten': '', 'huong_dan': '', 'quy_dinh': '', 'vien_dan': '', 'items': [], 'implicit': True}
                if cur_sec is None:
                    cur_sec = {'tt': '?', 'ten': '(khong xac dinh)', 'subs': [], 'huong_dan':'', 'quy_dinh':'', 'vien_dan':''}
                    sections.append(cur_sec)
                cur_sec['subs'].append(cur_sub)
            cur_sub['items'].append(item)
        else:
            if cur_sub:
                cur_sub['items'].append({'bullet': tt, 'ten': r['ten'], 'huong_dan': r['huong_dan'], 'quy_dinh': r['quy_dinh'], 'vien_dan': r['vien_dan']})

    # Post-process: ensure every section/sub with verifiable content has at least 1 item card.
    # If a sub-section has content (quy_dinh/huong_dan) but no items -> promote content into a single item.
    # If a section L1 has content but no subs -> create implicit sub + single item.
    for sec in sections:
        for sub in sec['subs']:
            if not sub['items'] and (sub.get('quy_dinh') or sub.get('huong_dan')):
                sub['items'].append({
                    '_row': sub.get('_row'),
                    'bullet': '◆',
                    'ten': sub['ten'] or '(mục này)',
                    'huong_dan': sub.get('huong_dan', ''),
                    'quy_dinh': sub.get('quy_dinh', ''),
                    'vien_dan': sub.get('vien_dan', '')
                })
                sub['huong_dan'] = ''; sub['quy_dinh'] = ''; sub['vien_dan'] = ''
        if not sec['subs']:
            sec['subs'].append({
                'tt': '', 'ten': '', 'implicit': True,
                'huong_dan': '', 'quy_dinh': '', 'vien_dan': '',
                'items': [{
                    '_row': sec.get('_row'),
                    'bullet': '◆',
                    'ten': sec['ten'] or '(mục này)',
                    'huong_dan': sec.get('huong_dan', ''),
                    'quy_dinh': sec.get('quy_dinh', ''),
                    'vien_dan': sec.get('vien_dan', '')
                }]
            })
            sec['huong_dan'] = ''; sec['quy_dinh'] = ''; sec['vien_dan'] = ''

    return sections, counter[0]

if __name__ == '__main__':
    fid = sys.argv[1] if len(sys.argv) > 1 else 'a4_van_phong'
    data, nimg = parse(fid)
    out_path = os.path.join(OUT, f'checklist_{fid}.js')
    js = 'window.BDC_CHECKLIST_' + fid.upper() + ' = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n'
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(js)
    nsec = len(data); nsub = sum(len(s['subs']) for s in data); nitem = sum(len(sub['items']) for s in data for sub in s['subs'])
    print(f'OK: {fid} -> {out_path}')
    print(f'  Sections L1={nsec}  Sub L2={nsub}  Items L3={nitem}  Images={nimg}')
